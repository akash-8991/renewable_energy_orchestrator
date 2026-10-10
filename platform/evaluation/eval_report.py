"""Excel report for one agent-evaluation run: what each check is, and how it went.

Pure data in -> .xlsx bytes out, with no agent or model imports, so the API can
build it. A stored run carries each case's outcome and (for newer runs) its
description; older runs are described from eval_catalogue. Text that flows from
the model or from an error message is neutralised against spreadsheet formula
injection, like the platform's other exports."""

from __future__ import annotations

import io
from datetime import datetime

try:  # as `evaluation.eval_report` (the API) or as a flat module (the agent worker's / tests' path layout)
    from .eval_catalogue import BY_NAME, TIER_MEANING
except ImportError:
    from eval_catalogue import BY_NAME, TIER_MEANING

FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")
STATUS_LABEL = {"passed": "PASSED", "failed": "FAILED", "skipped": "NOT SCORED"}
STATUS_FILL = {"PASSED": "C6EFCE", "FAILED": "FFC7CE", "NOT SCORED": "FFEB9C"}
STATUS_MEANING = [
    ("PASSED", "The agent met the check's expectation."),
    ("FAILED", "The agent ran but did not meet the expectation — or returned a response that failed schema validation after a retry (it fails closed)."),
    ("NOT SCORED", "The check was not counted: it is a reasoning check run on the mock provider, or the model was unavailable (billing, outage, rate limit, breaker). Says nothing about the agent."),
]


def _safe(value):
    if isinstance(value, str) and value and value[0] in FORMULA_TRIGGERS:
        return "'" + value
    return value


def _describe(result: dict) -> tuple[str, str]:
    """(description, tier meaning) for a stored result row."""
    info = BY_NAME.get(result.get("case", ""), {})
    description = result.get("description") or info.get("description") or "(no description recorded for this check)"
    return description, TIER_MEANING.get(result.get("tier") or info.get("tier", ""), "")


def build_eval_workbook(run: dict) -> bytes:
    """`run` has: id, created_at (datetime or ISO str), triggered_by, model_provider,
    total_cases, passed_cases, results (list of dicts)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    results = run.get("results") or []
    labels = [STATUS_LABEL.get(r.get("outcome", ""), str(r.get("outcome", "")).upper()) for r in results]
    passed, failed, skipped = labels.count("PASSED"), labels.count("FAILED"), labels.count("NOT SCORED")
    scored = passed + failed
    created = run.get("created_at")
    when = created.strftime("%Y-%m-%d %H:%M:%S UTC") if isinstance(created, datetime) else str(created)
    if scored == 0:
        verdict = "NOT SCORED"
    else:
        verdict = "ALL PASSED" if failed == 0 else f"{failed} FAILED"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    wrap = Alignment(wrap_text=True, vertical="top")

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.font, cell.fill, cell.alignment = header_font, header_fill, Alignment(vertical="center", wrap_text=True)

    wb = Workbook()

    # ---- Summary ----------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Agent evaluation run", ""])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    rows = [
        ("Run id", run.get("id", "")), ("Run time", when), ("Triggered by", run.get("triggered_by") or "system"),
        ("Model provider", run.get("model_provider", "")), ("Overall result", verdict),
        ("Checks scored", scored), ("Passed", passed), ("Failed", failed), ("Not scored", skipped),
        ("Pass rate (of scored)", f"{passed / scored:.0%}" if scored else "n/a"),
    ]
    for label, value in rows:
        ws.append([label, _safe(value)])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    ws.append(["What the statuses mean", ""])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for status, meaning in STATUS_MEANING:
        ws.append([status, meaning])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
        ws.cell(ws.max_row, 2).alignment = wrap
    ws.append([])
    ws.append(["What the check types mean", ""])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for tier, meaning in TIER_MEANING.items():
        ws.append([tier, meaning])
        ws.cell(ws.max_row, 2).alignment = wrap
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100

    # ---- Checks -----------------------------------------------------------------
    cs = wb.create_sheet("Checks")
    cs.append(["#", "Check", "Agent", "Type", "What it checks", "Status", "Result detail"])
    style_header(cs)
    for i, (r, label) in enumerate(zip(results, labels), start=1):
        description, _ = _describe(r)
        outcome = r.get("outcome")
        # for a passed check the stored detail only repeats the description — say it plainly instead
        detail = "Met the expectation." if outcome == "passed" else r.get("detail", "")
        cs.append([i, _safe(r.get("case", "")), _safe(r.get("agent", "")), _safe(r.get("tier", "")), _safe(description), label, _safe(detail)])
        status_cell = cs.cell(cs.max_row, 6)
        status_cell.fill = PatternFill("solid", fgColor=STATUS_FILL.get(label, "FFFFFF"))
        status_cell.font = Font(bold=True)
        for col in range(1, 8):
            cs.cell(cs.max_row, col).alignment = wrap
    for col, width in zip("ABCDEFG", (5, 46, 22, 12, 70, 14, 70)):
        cs.column_dimensions[col].width = width
    cs.freeze_panes = "A2"
    cs.auto_filter.ref = f"A1:{get_column_letter(7)}{max(cs.max_row, 2)}"

    # ---- By agent ---------------------------------------------------------------
    ag = wb.create_sheet("By agent")
    ag.append(["Agent", "Checks", "Passed", "Failed", "Not scored"])
    style_header(ag)
    per_agent: dict[str, list[str]] = {}
    for r, label in zip(results, labels):
        per_agent.setdefault(r.get("agent", ""), []).append(label)
    for agent in sorted(per_agent):
        statuses = per_agent[agent]
        ag.append([_safe(agent), len(statuses), statuses.count("PASSED"), statuses.count("FAILED"), statuses.count("NOT SCORED")])
    for col, width in zip("ABCDE", (28, 10, 10, 10, 12)):
        ag.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
