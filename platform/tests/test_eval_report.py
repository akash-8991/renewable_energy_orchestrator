"""The Excel report for an evaluation run: every check with what it tests and
its status. Must describe older runs (recorded before results carried a
description), label unscored checks honestly, and neutralise formula-injection
text; the catalogue it relies on must stay in step with the harness."""

import io
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent.parent / "evaluation"))
sys.path.insert(0, str(Path(__file__).parent.parent / "agent"))
sys.path.insert(0, str(Path(__file__).parent.parent / "backend"))

from app.routers import observability as obs  # noqa: E402
from database.connection import reset_current_tenant, set_current_tenant  # noqa: E402
from eval_catalogue import BY_NAME, CASES  # noqa: E402
from eval_harness import EVAL_CASES  # noqa: E402
from eval_report import build_eval_workbook  # noqa: E402
from fastapi import HTTPException  # noqa: E402
from models.canonical import AgentEvalRun, AuditEvent  # noqa: E402
from reo_common.security import AuthContext  # noqa: E402
from sqlalchemy import select  # noqa: E402

WHEN = datetime(2026, 10, 10, 17, 35, 10, tzinfo=timezone.utc)


def _run(results, **kw):
    return {"id": "run-1", "created_at": WHEN, "triggered_by": "m@test.example", "model_provider": "anthropic",
            "total_cases": kw.get("total", 0), "passed_cases": kw.get("passed", 0), "results": results}


def _sheets(content: bytes):
    wb = load_workbook(io.BytesIO(content))
    return wb, {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}


# --- the catalogue ----------------------------------------------------------------


def test_the_catalogue_matches_the_harness_exactly():
    assert [(c["name"], c["agent"], c["tier"], c["description"]) for c in CASES] == \
           [(c.name, c.agent, c.tier, c.description) for c in EVAL_CASES]


def test_every_result_the_harness_records_carries_its_description():
    from reo_common.model_gateway import MockModelGateway

    from eval_harness import run_eval_suite

    results = run_eval_suite(MockModelGateway())["results"]
    assert results and all(r.get("description") for r in results)
    assert {r["case"]: r["description"] for r in results} == {n: c["description"] for n, c in BY_NAME.items()}


# --- the workbook -------------------------------------------------------------------


def _first_behavioral(agent):
    return next(c["name"] for c in CASES if c["agent"] == agent and c["tier"] == "behavioral")


def test_the_workbook_has_a_summary_every_check_and_an_agent_rollup():
    asset_case, grid_case = _first_behavioral("asset"), _first_behavioral("grid")
    results = [
        {"case": "data_quality_schema_valid", "agent": "data_quality", "tier": "structural", "outcome": "passed",
         "description": BY_NAME["data_quality_schema_valid"]["description"], "detail": "x"},
        {"case": asset_case, "agent": "asset", "tier": "behavioral", "outcome": "failed", "detail": "agent failed closed: boom"},
        {"case": grid_case, "agent": "grid", "tier": "behavioral", "outcome": "skipped", "detail": "not scored — model unavailable"},
    ]
    wb, sheets = _sheets(build_eval_workbook(_run(results, total=2, passed=1)))
    assert wb.sheetnames == ["Summary", "Checks", "By agent"]

    summary = {row[0]: row[1] for row in sheets["Summary"] if row[0]}
    assert summary["Model provider"] == "anthropic" and summary["Checks scored"] == 2
    assert (summary["Passed"], summary["Failed"], summary["Not scored"]) == (1, 1, 1)
    assert summary["Overall result"] == "1 FAILED" and summary["Pass rate (of scored)"] == "50%"
    assert "PASSED" in summary and "NOT SCORED" in summary  # the legend is there

    checks = sheets["Checks"]
    assert checks[0] == ["#", "Check", "Agent", "Type", "What it checks", "Status", "Result detail"]
    by_check = {row[1]: row for row in checks[1:]}
    assert by_check["data_quality_schema_valid"][5] == "PASSED" and by_check["data_quality_schema_valid"][6] == "Met the expectation."
    # an older result with no stored description is still described, from the catalogue
    assert by_check[asset_case][4] == BY_NAME[asset_case]["description"]
    assert by_check[asset_case][5] == "FAILED" and "boom" in by_check[asset_case][6]
    assert by_check[grid_case][5] == "NOT SCORED"

    rollup = {row[0]: row[1:] for row in sheets["By agent"][1:]}
    assert rollup["asset"] == [1, 0, 1, 0] and rollup["grid"] == [1, 0, 0, 1]


def test_status_cells_are_colour_coded():
    results = [{"case": "a", "agent": "x", "tier": "structural", "outcome": o, "detail": ""} for o in ("passed", "failed", "skipped")]
    wb, _ = _sheets(build_eval_workbook(_run(results)))
    fills = [wb["Checks"].cell(row, 6).fill.fgColor.rgb[-6:] for row in (2, 3, 4)]
    assert fills == ["C6EFCE", "FFC7CE", "FFEB9C"]


def test_a_run_with_nothing_scored_says_so():
    results = [{"case": "a", "agent": "x", "tier": "behavioral", "outcome": "skipped", "detail": ""}]
    _, sheets = _sheets(build_eval_workbook(_run(results)))
    summary = {row[0]: row[1] for row in sheets["Summary"] if row[0]}
    assert summary["Overall result"] == "NOT SCORED" and summary["Pass rate (of scored)"] == "n/a"


def test_text_from_the_model_or_an_error_cannot_become_a_formula():
    results = [{"case": "=cmd|' /C calc'!A0", "agent": "@x", "tier": "structural", "outcome": "failed",
                "description": "+SUM(1,1)", "detail": "-2+3 boom"}]
    _, sheets = _sheets(build_eval_workbook(_run(results)))
    row = sheets["Checks"][1]
    assert all(str(v).startswith("'") for v in (row[1], row[2], row[4], row[6]))


# --- the endpoint -----------------------------------------------------------------------------


def _ctx(tenant_id):
    return AuthContext(user_id="00000000-0000-0000-0000-000000000001", tenant_id=tenant_id, roles=["viewer"], email="v@test.example")


def test_the_export_endpoint_returns_the_workbook_and_is_audited(db_session, two_tenants):
    t1, t2 = two_tenants
    token = set_current_tenant(t1.id)
    try:
        run = AgentEvalRun(tenant_id=t1.id, triggered_by="m@test.example", model_provider="anthropic", total_cases=1, passed_cases=1,
                           results=[{"case": "data_quality_schema_valid", "agent": "data_quality", "tier": "structural", "outcome": "passed", "detail": "d"}])
        db_session.add(run)
        db_session.flush()

        response = obs.export_eval_run(run.id, ctx=_ctx(t1.id), db=db_session)
        assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert 'attachment; filename="agent-evaluation-' in response.headers["Content-Disposition"]
        assert response.headers["Content-Disposition"].endswith('.xlsx"')
        _, sheets = _sheets(response.body)
        assert sheets["Checks"][1][1] == "data_quality_schema_valid"
        events = db_session.execute(select(AuditEvent.event_type).where(AuditEvent.tenant_id == t1.id)).scalars().all()
        assert "eval.exported" in events

        reset_current_tenant(token)
        token = set_current_tenant(t2.id)  # another tenant must not be able to download it
        with pytest.raises(HTTPException) as exc:
            obs.export_eval_run(run.id, ctx=_ctx(t2.id), db=db_session)
        assert exc.value.status_code == 404
    finally:
        reset_current_tenant(token)


def test_an_unknown_run_is_a_404(db_session, two_tenants):
    t1, _ = two_tenants
    token = set_current_tenant(t1.id)
    try:
        with pytest.raises(HTTPException) as exc:
            obs.export_eval_run(str(uuid.uuid4()), ctx=_ctx(t1.id), db=db_session)
        assert exc.value.status_code == 404
    finally:
        reset_current_tenant(token)
